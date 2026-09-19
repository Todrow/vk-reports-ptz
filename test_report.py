from datetime import date, datetime
from tempfile import TemporaryDirectory
from pathlib import Path
import unittest

from openpyxl import load_workbook

from database import Employee
from database import Database
from report import build_report
from bot import EMPLOYEES_COMMANDS, HELP_COMMANDS, REPORT_COMMANDS, attendance_summary, document_attachment, report_date_for_mark


class ReportTest(unittest.TestCase):
    def test_marks_after_noon_move_to_next_workday(self):
        workdays = frozenset({1, 2, 3, 4, 5})
        self.assertEqual(report_date_for_mark(datetime(2026, 8, 11, 11, 59), workdays), date(2026, 8, 11))
        self.assertEqual(report_date_for_mark(datetime(2026, 8, 11, 12, 0), workdays), date(2026, 8, 12))
        self.assertEqual(report_date_for_mark(datetime(2026, 8, 14, 18, 0), workdays), date(2026, 8, 17))

    def test_attendance_summary_excludes_automatic_unknown_statuses(self):
        self.assertEqual(attendance_summary({1: "present", 2: "unknown", 3: "unknown"}), (1, 2))

    def test_short_admin_commands_are_registered(self):
        self.assertIn("/р", REPORT_COMMANDS)
        self.assertIn("/с", EMPLOYEES_COMMANDS)
        self.assertIn("/к", HELP_COMMANDS)

    def test_document_attachment_accepts_vk_response_variants(self):
        document = {"owner_id": 1, "id": 2}
        self.assertEqual(document_attachment({"doc": document}), "doc1_2")
        self.assertEqual(document_attachment([document]), "doc1_2")

    def test_fills_template_and_keeps_single_status(self):
        employees = [
            Employee(1, "Петров Пётр Сергеевич", "1001", "24-37", True, True),
            Employee(2, "Сидорова Анна Игоревна", "1002", "24-37", True, True),
        ]
        with TemporaryDirectory() as directory:
            path = build_report(
                Path("Понедельник 10.08.2026.xlsx"),
                Path(directory),
                date(2026, 8, 11),
                "24-37",
                "Иванов Иван Иванович",
                employees,
                {1: "present", 2: "trip"},
            )
            ws = load_workbook(path).active
            self.assertEqual(ws["A1"].value, "Подразделение: 24-37")
            self.assertEqual(ws["A3"].value, "11.08.2026")
            self.assertEqual(ws["B5"].value, employees[0].full_name)
            self.assertEqual(ws["D5"].value, "+")
            self.assertEqual(ws["I6"].value, "+")
            self.assertEqual(ws.max_row, 6)
            self.assertEqual(path.name, "Вторник 11.08.2026.xlsx")

    def test_database_replaces_same_day_mark(self):
        from datetime import datetime

        with TemporaryDirectory() as directory:
            db = Database(Path(directory) / "test.sqlite3")
            db.register(42, "Петров Пётр Сергеевич", "1001", "24-37")
            db.approve(42)
            work_date = date(2026, 8, 11)
            db.mark(42, work_date, "present", datetime(2026, 8, 11, 8, 10))
            db.mark(42, work_date, "trip", datetime(2026, 8, 11, 8, 20))
            self.assertEqual(db.attendance(work_date), {42: "trip"})

    def test_unmarked_approved_employees_receive_unknown_status(self):
        from datetime import datetime

        with TemporaryDirectory() as directory:
            db = Database(Path(directory) / "test.sqlite3")
            db.register(1, "Петров Пётр Сергеевич", "1001", "24-37")
            db.register(2, "Сидорова Анна Игоревна", "1002", "24-37")
            db.approve(1)
            db.approve(2)
            work_date = date(2026, 8, 11)
            now = datetime(2026, 8, 11, 8, 30)
            db.mark(1, work_date, "present", now)
            db.mark_unmarked_as_unknown(work_date, now)
            self.assertEqual(db.attendance(work_date), {1: "present", 2: "unknown"})

    def test_manager_is_confirmed_separately_and_owns_one_bureau(self):
        with TemporaryDirectory() as directory:
            db = Database(Path(directory) / "test.sqlite3")
            candidate = db.request_manager(11, "Иванов Иван Иванович", "5001", "24-37")
            self.assertFalse(candidate.approved)
            manager = db.approve_manager(11)
            self.assertEqual(manager.bureau_number, "24-37")
            self.assertEqual(db.manager_for_bureau("24-37").vk_id, 11)
            db.register(12, "Петров Пётр Сергеевич", "5002", "24-37")
            self.assertEqual([person.vk_id for person in db.pending_employees("24-37")], [12])

    def test_employee_requires_approval_and_keeps_bureau_as_text(self):
        with TemporaryDirectory() as directory:
            db = Database(Path(directory) / "test.sqlite3")
            employee = db.register(42, "Петров Пётр Сергеевич", "1001", "24-37-А")
            self.assertEqual(employee.bureau_number, "24-37-А")
            self.assertFalse(employee.approved)
            self.assertEqual(db.employees(), [])
            db.approve(42)
            self.assertEqual(db.employees()[0].bureau_number, "24-37-А")
            db.reject(42)
            self.assertEqual(db.employees(), [])

    def test_lists_only_active_pending_employees(self):
        with TemporaryDirectory() as directory:
            db = Database(Path(directory) / "test.sqlite3")
            db.register(1, "Петров Пётр Сергеевич", "1001", "24-37")
            db.register(2, "Сидорова Анна Игоревна", "1002", "24-37")
            db.approve(2)
            self.assertEqual([employee.vk_id for employee in db.pending_employees()], [1])


if __name__ == "__main__":
    unittest.main()
