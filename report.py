from __future__ import annotations

from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from database import Employee


STATUS_COLUMNS = {
    "present": 4,
    "sick": 5,
    "day_off": 6,
    "unpaid": 7,
    "vacation": 8,
    "trip": 9,
    "scheduled_absence": 10,
    "unknown": 11,
}
FIRST_EMPLOYEE_ROW = 5
WEEKDAY_NAMES = (
    "Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье",
)


def _copy_row_style(ws, source_row: int, target_row: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, ws.max_column + 1):
        src, dst = ws.cell(source_row, col), ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)


def build_report(
    template: Path,
    output_dir: Path,
    work_date: date,
    bureau_number: str,
    boss_name: str,
    employees: list[Employee],
    attendance: dict[int, str],
) -> Path:
    wb = load_workbook(template)
    ws = wb.active
    ws.cell(1, 1, f"Подразделение: {bureau_number}")
    ws.cell(2, 1, f"Руководитель подразделения: {boss_name}")
    ws.cell(3, 1, work_date.strftime("%d.%m.%Y"))

    existing_rows = max(0, ws.max_row - FIRST_EMPLOYEE_ROW + 1)
    needed_rows = len(employees)
    if needed_rows > existing_rows:
        for row in range(FIRST_EMPLOYEE_ROW + existing_rows, FIRST_EMPLOYEE_ROW + needed_rows):
            _copy_row_style(ws, FIRST_EMPLOYEE_ROW, row)
    elif needed_rows < existing_rows:
        ws.delete_rows(FIRST_EMPLOYEE_ROW + needed_rows, existing_rows - needed_rows)

    for index, employee in enumerate(employees, start=1):
        row = FIRST_EMPLOYEE_ROW + index - 1
        for col in range(1, 12):
            ws.cell(row, col).value = None
        ws.cell(row, 1, index)
        ws.cell(row, 2, employee.full_name)
        ws.cell(row, 3, employee.personnel_number)
        status = attendance.get(employee.vk_id)
        if status in STATUS_COLUMNS:
            ws.cell(row, STATUS_COLUMNS[status], "+")

    output_dir.mkdir(parents=True, exist_ok=True)
    weekday = WEEKDAY_NAMES[work_date.weekday()]
    result = output_dir / f"{weekday} {work_date.strftime('%d.%m.%Y')}.xlsx"
    wb.save(result)
    return result
